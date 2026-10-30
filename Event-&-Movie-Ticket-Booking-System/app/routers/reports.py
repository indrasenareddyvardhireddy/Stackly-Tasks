from collections import defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth.dependencies import require_roles
from app.database import get_db
from app.models.booking import BookedSeat, Booking, Refund, Ticket
from app.models.enums import BookingStatus, EventStatus, RefundReason, Role
from app.models.event import Event, Show
from app.models.user import User
from app.models.venue import Seat
from app.schemas.report import (AdminDashboard, CategoryOccupancy, EventSales, RefundReport, RefundSummaryRow,
                                ShowOccupancy, ShowSales, TopEvent)
from app.utils.billing import money
from app.utils.time import now

router = APIRouter(prefix="/reports", tags=["Reports & Dashboard"])
admin_only = require_roles(Role.ADMIN)
admin_or_organizer = require_roles(Role.ADMIN, Role.ORGANIZER)
CONFIRMED = BookingStatus.CONFIRMED.value


def _day_bounds(d: date):
    start = datetime.combine(d, time.min)
    return start, start + timedelta(days=1)


def _scoped_shows(db: Session, user: User, from_date: Optional[date], to_date: Optional[date],
                  event_id: Optional[int] = None):
    q = db.query(Show).join(Event, Event.id == Show.event_id)
    if user.role == Role.ORGANIZER.value:
        q = q.filter(Event.organizer_id == user.id)
    if event_id:
        q = q.filter(Event.id == event_id)
    if from_date:
        q = q.filter(Show.show_date >= from_date)
    if to_date:
        q = q.filter(Show.show_date <= to_date)
    return q.order_by(Show.start_at)


@router.get("/admin-dashboard", response_model=AdminDashboard)
def admin_dashboard(db: Session = Depends(get_db), _: User = Depends(admin_only)):
    start, end = _day_bounds(now().date())
    bookings_today, revenue_today = db.query(
        func.count(Booking.id), func.coalesce(func.sum(Booking.grand_total), 0)
    ).filter(Booking.confirmed_at >= start, Booking.confirmed_at < end).one()
    refund_count, refund_amount = db.query(
        func.count(Refund.id), func.coalesce(func.sum(Refund.amount), 0)
    ).filter(Refund.refunded_at >= start, Refund.refunded_at < end).one()
    active_events = db.query(func.count(Event.id)).filter(Event.status == EventStatus.PUBLISHED.value).scalar()
    scanned = db.query(func.count(Ticket.id)).filter(Ticket.used_at >= start, Ticket.used_at < end).scalar()
    return AdminDashboard(date=start.date(), todays_bookings=bookings_today, todays_revenue=money(revenue_today),
                          todays_refund_count=refund_count, todays_refunds=money(refund_amount),
                          active_events=active_events, tickets_scanned_today=scanned)


def event_sales_data(db: Session, user: User, from_date, to_date, event_id=None) -> List[EventSales]:
    shows = _scoped_shows(db, user, from_date, to_date, event_id).all()
    if not shows:
        return []
    ids = [s.id for s in shows]
    sold = {sid: (int(n), Decimal(rev)) for sid, n, rev in db.query(
        Booking.show_id, func.coalesce(func.sum(Booking.seat_count), 0), func.coalesce(func.sum(Booking.grand_total), 0)
    ).filter(Booking.show_id.in_(ids), Booking.status == CONFIRMED).group_by(Booking.show_id).all()}
    refunds = {sid: Decimal(amt) for sid, amt in db.query(Booking.show_id, func.coalesce(func.sum(Refund.amount), 0))
               .join(Refund, Refund.booking_id == Booking.id).filter(Booking.show_id.in_(ids))
               .group_by(Booking.show_id).all()}

    events: Dict[int, EventSales] = {}
    for s in shows:
        tickets, revenue = sold.get(s.id, (0, Decimal("0")))
        refunded = refunds.get(s.id, Decimal("0"))
        ev = events.setdefault(s.event_id, EventSales(event_id=s.event_id, title=s.event.title, tickets_sold=0,
                                                      revenue=Decimal("0"), refunds=Decimal("0"), shows=[]))
        ev.tickets_sold += tickets
        ev.revenue += revenue
        ev.refunds += refunded
        ev.shows.append(ShowSales(show_id=s.id, show_start=s.start_at.strftime("%Y-%m-%d %H:%M"),
                                  venue_name=s.venue_name, tickets_sold=tickets, revenue=revenue, refunds=refunded))
    return list(events.values())


@router.get("/event-sales", response_model=List[EventSales])
def event_sales(from_date: Optional[date] = None, to_date: Optional[date] = None, event_id: Optional[int] = None,
                db: Session = Depends(get_db), user: User = Depends(admin_or_organizer)):
    return event_sales_data(db, user, from_date, to_date, event_id)


@router.get("/event-sales/export")
def event_sales_excel(from_date: Optional[date] = None, to_date: Optional[date] = None,
                      event_id: Optional[int] = None, db: Session = Depends(get_db),
                      user: User = Depends(admin_or_organizer)):
    """Bonus: Excel export of the sales report."""
    data = event_sales_data(db, user, from_date, to_date, event_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Event Sales"
    ws.append(["Event", "Show ID", "Show Start", "Venue", "Tickets Sold", "Revenue (Rs.)", "Refunds (Rs.)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for ev in data:
        for sh in ev.shows:
            ws.append([ev.title, sh.show_id, sh.show_start, sh.venue_name, sh.tickets_sold,
                       float(sh.revenue), float(sh.refunds)])
        ws.append([f"TOTAL - {ev.title}", "", "", "", ev.tickets_sold, float(ev.revenue), float(ev.refunds)])
        ws[ws.max_row][0].font = Font(bold=True)
    for col, width in zip("ABCDEFG", (34, 10, 18, 28, 14, 16, 16)):
        ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=event_sales.xlsx"})


@router.get("/occupancy", response_model=List[ShowOccupancy])
def occupancy(from_date: Optional[date] = None, to_date: Optional[date] = None, event_id: Optional[int] = None,
              db: Session = Depends(get_db), user: User = Depends(admin_or_organizer)):
    shows = _scoped_shows(db, user, from_date, to_date, event_id).limit(500).all()
    if not shows:
        return []
    screen_ids = {s.screen_id for s in shows}
    totals: Dict[int, Dict[str, int]] = defaultdict(dict)
    for screen_id, category, n in (db.query(Seat.screen_id, Seat.category, func.count(Seat.id))
                                   .filter(Seat.screen_id.in_(screen_ids), Seat.is_active.is_(True))
                                   .group_by(Seat.screen_id, Seat.category).all()):
        totals[screen_id][category] = n
    booked: Dict[int, Dict[str, int]] = defaultdict(dict)
    for show_id, category, n in (db.query(BookedSeat.show_id, Seat.category, func.count(BookedSeat.id))
                                 .join(Seat, Seat.id == BookedSeat.seat_id)
                                 .filter(BookedSeat.show_id.in_([s.id for s in shows]),
                                         BookedSeat.confirmed_seat_id.isnot(None))
                                 .group_by(BookedSeat.show_id, Seat.category).all()):
        booked[show_id][category] = n

    out = []
    for s in shows:
        cats, total, taken = [], 0, 0
        for category, cat_total in sorted(totals[s.screen_id].items()):
            cat_booked = booked[s.id].get(category, 0)
            total += cat_total
            taken += cat_booked
            cats.append(CategoryOccupancy(category=category, total_seats=cat_total, booked_seats=cat_booked,
                                          occupancy_percent=round(cat_booked * 100 / cat_total, 2)))
        out.append(ShowOccupancy(show_id=s.id, event_title=s.event.title,
                                 show_start=s.start_at.strftime("%Y-%m-%d %H:%M"), total_seats=total,
                                 booked_seats=taken, occupancy_percent=round(taken * 100 / total, 2) if total else 0.0,
                                 categories=cats))
    return out


@router.get("/top-events", response_model=List[TopEvent])
def top_events(from_date: Optional[date] = None, to_date: Optional[date] = None,
               db: Session = Depends(get_db), _: User = Depends(admin_only)):
    q = (db.query(Event.id, Event.title, func.sum(Booking.seat_count).label("tickets"),
                  func.sum(Booking.grand_total).label("revenue"))
         .join(Show, Show.event_id == Event.id).join(Booking, Booking.show_id == Show.id)
         .filter(Booking.status == CONFIRMED))
    if from_date:
        q = q.filter(Show.show_date >= from_date)
    if to_date:
        q = q.filter(Show.show_date <= to_date)
    rows = q.group_by(Event.id, Event.title).order_by(func.sum(Booking.seat_count).desc()).limit(10).all()
    return [TopEvent(event_id=i, title=t, tickets_sold=int(n), revenue=money(rev)) for i, t, n, rev in rows]


@router.get("/refunds", response_model=RefundReport)
def refund_report(from_date: Optional[date] = None, to_date: Optional[date] = None,
                  db: Session = Depends(get_db), _: User = Depends(admin_only)):
    q = db.query(Refund.reason, func.count(Refund.id), func.coalesce(func.sum(Refund.amount), 0))
    if from_date:
        q = q.filter(Refund.refunded_at >= datetime.combine(from_date, time.min))
    if to_date:
        q = q.filter(Refund.refunded_at <= datetime.combine(to_date, time.max))
    found = {reason: (n, Decimal(amount)) for reason, n, amount in q.group_by(Refund.reason).all()}
    rows = [RefundSummaryRow(reason=r.value, count=found.get(r.value, (0, Decimal("0")))[0],
                             amount=found.get(r.value, (0, Decimal("0")))[1]) for r in RefundReason]
    return RefundReport(total_count=sum(r.count for r in rows), total_amount=sum((r.amount for r in rows), Decimal("0")),
                        by_reason=rows)
